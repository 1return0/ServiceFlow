from __future__ import annotations

from pathlib import Path
from zipfile import BadZipFile

from pptx import Presentation
from pptx.exc import PackageNotFoundError

from app.domain.parsed_support_knowledge_document import (
    ParsedSupportKnowledgeDocument,
    ParsedSupportKnowledgeFragment,
    SupportKnowledgeArtifactParseError,
)
from app.domain.support_knowledge_artifact import SupportKnowledgeArtifact


class PptxSupportKnowledgeArtifactParser:
    """Extract visible text frames in slide shape order.

    This deliberately excludes speaker notes, charts, and text embedded in images.
    """

    def __init__(self, artifact_root: Path) -> None:
        self._artifact_root = artifact_root.resolve()

    @property
    def document_format(self) -> str:
        return "pptx"

    def parse(
        self,
        artifact: SupportKnowledgeArtifact,
    ) -> ParsedSupportKnowledgeDocument:
        if artifact.document_format != self.document_format:
            raise SupportKnowledgeArtifactParseError(
                "PPTX parser received a non-PPTX artifact"
            )
        path = self._resolve_artifact_path(artifact.storage_relative_path)
        try:
            presentation = Presentation(path)
        except (OSError, BadZipFile, PackageNotFoundError, ValueError) as error:
            raise SupportKnowledgeArtifactParseError(
                "PPTX artifact could not be read"
            ) from error

        fragments = _extract_fragments(presentation)
        if not fragments:
            raise SupportKnowledgeArtifactParseError(
                "PPTX artifact contains no extractable text"
            )
        return ParsedSupportKnowledgeDocument(
            artifact_id=artifact.artifact_id,
            document_format="pptx",
            title=None,
            text="\n".join(fragment.text for fragment in fragments),
            fragments=fragments,
        )

    def _resolve_artifact_path(self, storage_relative_path: str) -> Path:
        path = (self._artifact_root / storage_relative_path).resolve()
        if self._artifact_root not in path.parents:
            raise SupportKnowledgeArtifactParseError(
                "artifact path escaped the configured artifact root"
            )
        return path


def _extract_fragments(presentation) -> tuple[ParsedSupportKnowledgeFragment, ...]:
    fragments: list[ParsedSupportKnowledgeFragment] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        for shape_number, shape in enumerate(slide.shapes, start=1):
            if not shape.has_text_frame:
                continue
            text = "\n".join(
                paragraph.text.strip()
                for paragraph in shape.text_frame.paragraphs
                if paragraph.text.strip()
            )
            if not text:
                continue
            fragments.append(
                ParsedSupportKnowledgeFragment(
                    ordinal=len(fragments) + 1,
                    location=f"slide:{slide_number}:shape:{shape_number}",
                    text=text,
                )
            )
    return tuple(fragments)
