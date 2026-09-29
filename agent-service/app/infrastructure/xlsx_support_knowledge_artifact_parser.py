from __future__ import annotations

from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.domain.parsed_support_knowledge_document import (
    ParsedSupportKnowledgeDocument,
    ParsedSupportKnowledgeFragment,
    SupportKnowledgeArtifactParseError,
)
from app.domain.support_knowledge_artifact import SupportKnowledgeArtifact


class XlsxSupportKnowledgeArtifactParser:
    """Extract non-empty worksheet rows without inventing formula results."""

    def __init__(self, artifact_root: Path) -> None:
        self._artifact_root = artifact_root.resolve()

    @property
    def document_format(self) -> str:
        return "xlsx"

    def parse(
        self,
        artifact: SupportKnowledgeArtifact,
    ) -> ParsedSupportKnowledgeDocument:
        if artifact.document_format != self.document_format:
            raise SupportKnowledgeArtifactParseError(
                "XLSX parser received a non-XLSX artifact"
            )
        path = self._resolve_artifact_path(artifact.storage_relative_path)
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except (OSError, BadZipFile, InvalidFileException, ValueError) as error:
            raise SupportKnowledgeArtifactParseError(
                "XLSX artifact could not be read"
            ) from error

        try:
            fragments = _extract_fragments(workbook)
        finally:
            workbook.close()
        if not fragments:
            raise SupportKnowledgeArtifactParseError(
                "XLSX artifact contains no extractable text"
            )
        return ParsedSupportKnowledgeDocument(
            artifact_id=artifact.artifact_id,
            document_format="xlsx",
            title=None,
            text="\n".join(fragment.text for fragment in fragments),
            fragments=fragments,
        )

    def _resolve_artifact_path(self, storage_relative_path: str) -> Path:
        path = (self._artifact_root / storage_relative_path).resolve()
        if self._artifact_root not in path.parents:
            raise SupportKnowledgeArtifactParseError(
                "artifact path escaped the configured artifact root"
            )
        return path


def _extract_fragments(workbook) -> tuple[ParsedSupportKnowledgeFragment, ...]:
    fragments: list[ParsedSupportKnowledgeFragment] = []
    for worksheet in workbook.worksheets:
        for row_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
            cells = tuple(_cell_text(value) for value in row)
            non_blank_cells = tuple(cell for cell in cells if cell)
            if not non_blank_cells:
                continue
            fragments.append(
                ParsedSupportKnowledgeFragment(
                    ordinal=len(fragments) + 1,
                    location=f"sheet:{worksheet.title}:row:{row_number}",
                    text=" | ".join(non_blank_cells),
                )
            )
    return tuple(fragments)


def _cell_text(value: object | None) -> str:
    if value is None:
        return ""
    return str(value).strip()
