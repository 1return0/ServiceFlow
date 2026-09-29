from __future__ import annotations

from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase

from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from pptx import Presentation

from app.bootstrap import build_support_knowledge_document_parsing_service
from app.domain.support_knowledge_artifact import SupportKnowledgeArtifact


class SupportKnowledgeDocumentParsingServiceTest(TestCase):
    def test_routes_five_real_artifact_formats_from_one_composition_root(self) -> None:
        with TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "guide.html").write_text(
                "<title>HTML Guide</title><p>Network client instructions</p>",
                encoding="utf-8",
            )
            _write_pdf_with_text(root / "guide.pdf", "PDF network instructions")
            docx = Document()
            docx.add_paragraph("DOCX network instructions")
            docx.save(root / "guide.docx")
            workbook = Workbook()
            workbook.active.append(["XLSX network instructions"])
            workbook.save(root / "guide.xlsx")
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            slide.shapes.add_textbox(0, 0, 1000000, 300000).text = (
                "PPTX network instructions"
            )
            presentation.save(root / "guide.pptx")
            service = build_support_knowledge_document_parsing_service(root)

            html = service.parse(_artifact("html", "guide.html"))
            pdf = service.parse(_artifact("pdf", "guide.pdf"))
            docx_parsed = service.parse(_artifact("docx", "guide.docx"))
            xlsx = service.parse(_artifact("xlsx", "guide.xlsx"))
            pptx = service.parse(_artifact("pptx", "guide.pptx"))

        self.assertEqual("HTML Guide", html.title)
        self.assertEqual("HTML Guide Network client instructions", html.text)
        self.assertEqual("PDF network instructions", pdf.text)
        self.assertEqual("DOCX network instructions", docx_parsed.text)
        self.assertEqual("XLSX network instructions", xlsx.text)
        self.assertEqual("PPTX network instructions", pptx.text)


def _write_pdf_with_text(path: Path, text: str) -> None:
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
    )
    content = DecodedStreamObject()
    content.set_data(f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("ascii"))
    page[NameObject("/Contents")] = writer._add_object(content)
    with path.open("wb") as stream:
        writer.write(stream)


def _artifact(document_format: str, storage_relative_path: str) -> SupportKnowledgeArtifact:
    return SupportKnowledgeArtifact(
        artifact_id=f"artifact_{document_format}",
        source_reference="https://it.example.edu/guide",
        source_kind="official_campus",
        document_format=document_format,  # type: ignore[arg-type]
        storage_relative_path=storage_relative_path,
        byte_size=1,
        content_sha256="a" * 64,
    )
