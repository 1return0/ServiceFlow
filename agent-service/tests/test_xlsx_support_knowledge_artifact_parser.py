from __future__ import annotations

from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase

from openpyxl import Workbook

from app.domain.parsed_support_knowledge_document import (
    SupportKnowledgeArtifactParseError,
)
from app.domain.support_knowledge_artifact import SupportKnowledgeArtifact
from app.infrastructure.xlsx_support_knowledge_artifact_parser import (
    XlsxSupportKnowledgeArtifactParser,
)


class XlsxSupportKnowledgeArtifactParserTest(TestCase):
    def test_extracts_non_empty_rows_with_sheet_and_row_locations(self) -> None:
        with TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "service-hours.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "IT服务"
            worksheet.append(["服务", "地点"])
            worksheet.append(["网络支持", "教学楼 A"])
            worksheet.append([None, None])
            workbook.create_sheet("图书馆").append(["借阅", "一楼服务台"])
            workbook.save(path)

            parsed = XlsxSupportKnowledgeArtifactParser(root).parse(
                _artifact(path, "service-hours.xlsx")
            )

            self.assertEqual(
                "服务 | 地点\n网络支持 | 教学楼 A\n借阅 | 一楼服务台",
                parsed.text,
            )
            self.assertEqual(
                ("sheet:IT服务:row:1", "sheet:IT服务:row:2", "sheet:图书馆:row:1"),
                tuple(fragment.location for fragment in parsed.fragments),
            )

    def test_rejects_a_workbook_without_text(self) -> None:
        with TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "empty.xlsx"
            Workbook().save(path)

            with self.assertRaises(SupportKnowledgeArtifactParseError):
                XlsxSupportKnowledgeArtifactParser(root).parse(
                    _artifact(path, "empty.xlsx")
                )


def _artifact(path: Path, relative_path: str) -> SupportKnowledgeArtifact:
    return SupportKnowledgeArtifact(
        artifact_id="artifact-xlsx",
        source_reference="https://it.example.edu/service-hours.xlsx",
        source_kind="official_campus",
        document_format="xlsx",
        storage_relative_path=relative_path,
        byte_size=path.stat().st_size,
        content_sha256="xlsx-sha",
    )
