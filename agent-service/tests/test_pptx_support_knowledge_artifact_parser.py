from __future__ import annotations

from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase

from pptx import Presentation

from app.domain.parsed_support_knowledge_document import (
    SupportKnowledgeArtifactParseError,
)
from app.domain.support_knowledge_artifact import SupportKnowledgeArtifact
from app.infrastructure.parsers.pptx_support_knowledge_artifact_parser import (
    PptxSupportKnowledgeArtifactParser,
)


class PptxSupportKnowledgeArtifactParserTest(TestCase):
    def test_extracts_text_frames_in_slide_and_shape_order(self) -> None:
        with TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "network-guide.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            slide.shapes.add_textbox(0, 0, 1000000, 300000).text = "校园网络指南"
            slide.shapes.add_textbox(0, 400000, 1000000, 300000).text = "先连接 CampusNet"
            presentation.save(path)

            parsed = PptxSupportKnowledgeArtifactParser(root).parse(
                _artifact(path, "network-guide.pptx")
            )

            self.assertEqual("校园网络指南\n先连接 CampusNet", parsed.text)
            self.assertEqual(
                ("slide:1:shape:1", "slide:1:shape:2"),
                tuple(fragment.location for fragment in parsed.fragments),
            )

    def test_rejects_a_presentation_without_text(self) -> None:
        with TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "empty.pptx"
            Presentation().save(path)

            with self.assertRaises(SupportKnowledgeArtifactParseError):
                PptxSupportKnowledgeArtifactParser(root).parse(
                    _artifact(path, "empty.pptx")
                )


def _artifact(path: Path, relative_path: str) -> SupportKnowledgeArtifact:
    return SupportKnowledgeArtifact(
        artifact_id="artifact-pptx",
        source_reference="https://it.example.edu/network-guide.pptx",
        source_kind="official_campus",
        document_format="pptx",
        storage_relative_path=relative_path,
        byte_size=path.stat().st_size,
        content_sha256="pptx-sha",
    )
