from __future__ import annotations

from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase

from docx import Document

from app.domain.parsed_support_knowledge_document import (
    SupportKnowledgeArtifactParseError,
)
from app.domain.support_knowledge_artifact import SupportKnowledgeArtifact
from app.infrastructure.parsers.docx_support_knowledge_artifact_parser import (
    DocxSupportKnowledgeArtifactParser,
)


class DocxSupportKnowledgeArtifactParserTest(TestCase):
    def test_extracts_core_title_paragraphs_and_table_rows(self) -> None:
        with TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "network-guide.docx"
            document = Document()
            document.core_properties.title = "Campus network guide"
            document.add_paragraph("First connect to CampusNet.")
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "Service desk"
            table.cell(0, 1).text = "Building A"
            document.add_paragraph("Then complete web authentication.")
            document.save(path)

            parsed = DocxSupportKnowledgeArtifactParser(root).parse(
                SupportKnowledgeArtifact(
                    artifact_id="artifact-docx-1",
                    source_reference="https://it.example.edu/network-guide.docx",
                    source_kind="official_campus",
                    document_format="docx",
                    storage_relative_path="network-guide.docx",
                    byte_size=path.stat().st_size,
                    content_sha256="docx-sha",
                )
            )

            self.assertEqual(parsed.title, "Campus network guide")
            self.assertEqual(
                parsed.text,
                "First connect to CampusNet.\n"
                "Service desk | Building A\n"
                "Then complete web authentication.",
            )

    def test_rejects_a_docx_with_no_extractable_text(self) -> None:
        with TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "empty.docx"
            Document().save(path)

            with self.assertRaises(SupportKnowledgeArtifactParseError):
                DocxSupportKnowledgeArtifactParser(root).parse(
                    SupportKnowledgeArtifact(
                        artifact_id="artifact-docx-empty",
                        source_reference="https://it.example.edu/empty.docx",
                        source_kind="official_campus",
                        document_format="docx",
                        storage_relative_path="empty.docx",
                        byte_size=path.stat().st_size,
                        content_sha256="empty-docx-sha",
                    )
                )
